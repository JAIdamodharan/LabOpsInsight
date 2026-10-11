import datetime as dt
import io
import zipfile

import duckdb
import pytest
from openpyxl import Workbook

from analyst import ingest
from analyst.ingest import IngestError, ident


def xlsx(sheets: dict, hidden=()):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        if name in hidden:
            ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def load(tmp_path, data, name="book.xlsx"):
    res = ingest.ingest(name, data, tmp_path / "d.duckdb")
    con = duckdb.connect(str(tmp_path / "d.duckdb"), read_only=True)
    return res, con


def test_types_dates_nulls_and_header_cleaning(tmp_path):
    data = xlsx({"Sales 2025": [
        ["Order Date", "Region", "Units", "Unit Price", "Paid?", "Notes", None, "Order Date"],
        [dt.datetime(2025, 1, 5), "North", 3, 9.5, True, None, None, "x"],
        [dt.datetime(2025, 1, 6), "South", 4, 10.0, False, "late", None, "y"],
        [None, "North", None, 8.25, True, None, None, "z"]]})
    res, con = load(tmp_path, data)
    t = res.tables[0]
    assert t.name == "sales_2025" and t.sheet == "Sales 2025" and t.rows == 3
    cols = {c["name"]: c for c in t.columns}
    assert list(cols) == ["order_date", "region", "units", "unit_price", "paid", "notes", "order_date_2"]   # blank 7th column dropped, duplicate header suffixed
    types = {c["name"]: c["type"] for c in t.columns}
    assert types["order_date"] == "DATE" and types["units"] == "BIGINT" and types["unit_price"] == "DOUBLE" and types["paid"] == "BOOLEAN"
    assert cols["order_date"]["original"] == "Order Date"
    assert con.execute("SELECT count(*) FROM sales_2025 WHERE units IS NULL").fetchone()[0] == 1
    assert con.execute("SELECT sum(units) FROM sales_2025").fetchone()[0] == 7


def test_empty_trailing_column_dropped(tmp_path):
    res, _ = load(tmp_path, xlsx({"S": [["a", "b", None], [1, 2, None], [3, 4, None]]}))
    assert [c["name"] for c in res.tables[0].columns] == ["a", "b"]


def test_numbers_stored_as_text_are_kept_as_text(tmp_path):
    res, con = load(tmp_path, xlsx({"S": [["amount"], ["100"], ["250"], ["N/A"]]}))
    assert res.tables[0].columns[0]["type"] == "VARCHAR"
    assert con.execute("SELECT amount FROM s ORDER BY 1 LIMIT 1").fetchone()[0] == "100"


def test_mixed_types_become_text_with_warning(tmp_path):
    res, con = load(tmp_path, xlsx({"S": [["v"], [1], ["two"], [3.5]]}))
    assert res.tables[0].columns[0]["type"] == "VARCHAR" and any("mixes text" in w for w in res.warnings)


def test_multiple_sheets_hidden_and_empty_sheets(tmp_path):
    data = xlsx({"Orders": [["id"], [1], [2]], "Secret": [["x"], [1]], "Empty": [], "Header only": [["a", "b"]]}, hidden={"Secret"})
    res, _ = load(tmp_path, data)
    assert [t.name for t in res.tables] == ["orders"]
    assert any("Hidden sheet 'Secret'" in w for w in res.warnings) and any("no data rows" in w for w in res.warnings)


def test_reserved_and_duplicate_names():
    used = set()
    assert ident("Order", "c", used) == "order_" and ident("order", "c", used) == "order__2"
    u2 = set()
    assert [ident(x, "column", u2) for x in ["Total $", "total", "TOTAL", "2024"]] == ["total", "total_2", "total_3", "column_2024"]
    assert ident("   ", "column_1", set()) == "column_1" and len(ident("x" * 200, "c", set())) <= 60


def test_csv(tmp_path):
    res, con = load(tmp_path, b"Name,Score,When\nAnn,3.5,2025-01-01\nBob,,2025-01-02\n", "My Data.csv")
    assert res.tables[0].name == "my_data" and res.tables[0].rows == 2
    assert con.execute("SELECT count(*) FROM my_data WHERE score IS NULL").fetchone()[0] == 1


def test_csv_latin1_fallback(tmp_path):
    res, con = load(tmp_path, "city,n\nZürich,1\n".encode("latin-1"), "c.csv")
    assert con.execute("SELECT city FROM c").fetchone()[0] == "Zürich"


@pytest.mark.parametrize("name,data,msg", [
    ("a.xls", b"x", "Legacy .xls"),
    ("a.xlsm", b"x", "Macro-enabled"),
    ("a.txt", b"x", "Only .xlsx and .csv"),
    ("a.xlsx", b"", "empty"),
    ("a.xlsx", b"not a zip at all", "not a valid .xlsx"),
    ("a.xlsx", bytes.fromhex("D0CF11E0A1B11AE1") + b"0" * 50, "legacy .xls file or a password-protected"),
    ("a.csv", b"a,b\n", "no data rows"),
])
def test_rejections(tmp_path, name, data, msg):
    with pytest.raises(IngestError, match=msg):
        ingest.ingest(name, data, tmp_path / "d.duckdb")


def test_zip_that_is_not_a_workbook(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(IngestError, match="not an Excel workbook"):
        ingest.ingest("a.xlsx", buf.getvalue(), tmp_path / "d.duckdb")


def test_macro_workbook_rejected(tmp_path):
    good = xlsx({"S": [["a"], [1]]})
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(good)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for i in zin.infolist():
            zout.writestr(i, zin.read(i.filename))
        zout.writestr("xl/vbaProject.bin", b"fake")
    with pytest.raises(IngestError, match="Macro-enabled"):
        ingest.ingest("a.xlsx", buf.getvalue(), tmp_path / "d.duckdb")


def test_zip_bomb_rejected_by_expansion_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "MAX_UNCOMPRESSED", 10_000)
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx({"S": [["a"], [1]]}))) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            zout.writestr(i, zin.read(i.filename))
        zout.writestr("xl/bomb.xml", b"0" * 5_000_000)           # compresses to a few KB, expands to 5 MB
    assert len(buf.getvalue()) < 100_000
    with pytest.raises(IngestError, match="expands to more than"):
        ingest.ingest("a.xlsx", buf.getvalue(), tmp_path / "d.duckdb")


def test_size_row_col_and_sheet_limits(tmp_path, monkeypatch):
    with pytest.raises(IngestError, match="limit is 15 MB"):
        ingest.ingest("a.csv", b"a\n" + b"1\n" * 8_000_000, tmp_path / "d.duckdb")
    monkeypatch.setattr(ingest, "MAX_ROWS", 5)
    with pytest.raises(IngestError, match="more than 5 rows"):
        ingest.ingest("a.xlsx", xlsx({"S": [["a"]] + [[i] for i in range(10)]}), tmp_path / "d.duckdb")
    with pytest.raises(IngestError, match="more than 5 rows"):
        ingest.ingest("a.csv", b"a\n" + b"1\n" * 10, tmp_path / "d.duckdb")
    monkeypatch.setattr(ingest, "MAX_ROWS", 200_000)
    monkeypatch.setattr(ingest, "MAX_COLS", 3)
    with pytest.raises(IngestError, match="4 columns"):
        ingest.ingest("a.xlsx", xlsx({"S": [["a", "b", "c", "d"], [1, 2, 3, 4]]}), tmp_path / "d.duckdb")
    monkeypatch.setattr(ingest, "MAX_COLS", 100)
    monkeypatch.setattr(ingest, "MAX_SHEETS", 2)
    with pytest.raises(IngestError, match="more than 2 sheets"):
        ingest.ingest("a.xlsx", xlsx({f"S{i}": [["a"], [1]] for i in range(3)}), tmp_path / "d.duckdb")
    monkeypatch.setattr(ingest, "MAX_SHEETS", 10)
    monkeypatch.setattr(ingest, "MAX_CELLS", 10)
    with pytest.raises(IngestError, match="cells"):
        ingest.ingest("a.xlsx", xlsx({"S": [["a", "b"]] + [[i, i] for i in range(20)]}), tmp_path / "d.duckdb")


def test_formula_cells_use_cached_values_and_are_never_evaluated(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["a", "b"])
    ws.append([1, "=1+1"])          # no cached value because it was never calculated by Excel
    buf = io.BytesIO()
    wb.save(buf)
    res, con = load(tmp_path, buf.getvalue())
    assert con.execute("SELECT b FROM s").fetchone()[0] is None


def test_hostile_cell_text_is_just_data(tmp_path):
    evil = "'; DROP TABLE s; --"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["note"])
    ws.append([evil])
    ws.append(["x"])
    ws["A4"] = "=cmd|' /C calc'!A0"
    ws["A4"].data_type = "s"                      # a text cell that merely *looks* like a formula
    buf = io.BytesIO()
    wb.save(buf)
    res, con = load(tmp_path, buf.getvalue())
    assert con.execute("SELECT count(*) FROM s").fetchone()[0] == 3
    assert {r[0] for r in con.execute("SELECT note FROM s").fetchall()} == {evil, "x", "=cmd|' /C calc'!A0"}
    assert con.execute("SELECT count(*) FROM information_schema.tables WHERE table_name = 's'").fetchone()[0] == 1
